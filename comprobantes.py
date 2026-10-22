"""
comprobantes.py
===============
Impresión (PDF carta y Excel) de los asientos contables de la planilla:

* ``pdf_asiento_privado`` / ``excel_asiento_privado``: asiento para empresa privada.
* ``pdf_comprobante_publico`` / ``excel_comprobante_publico``: comprobante
  patrimonial para el sector público.

Los datos vienen de ``asientos.py``. ``meta`` es el mismo diccionario que usa
``documentos.py`` y, para el sector público, además ``meta["comprobante"]``.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

from documentos import (COLOR_HEADER, COLOR_TOTAL, _firmas_pdf, _fmt_fecha, _fmt_monto,
                        _pie)

NEGRO = colors.black
GRIS = colors.HexColor("#" + COLOR_TOTAL)

COMPROBANTE_POR_DEFECTO: dict[str, Any] = {
    "numero": "",
    "areas": ["Presupuesto", "Patrimonio"],
    "imputacion": "CON IMPUTACIÓN PRESUPUESTARIA",
    "operacion": None,            # Balance Inicial / Traspaso Caja-Banco / Ajuste Contable
    "momento_recursos": None,     # Estimación / Devengado / Percibido
    "momento_gastos": "DEVENGADO",  # Apropiación / Preventivo / Compromiso / Devengado / Pagado
    "beneficiario": "",
    "respaldo": "Planilla de sueldos y salarios",
    "apertura": "---/---/---", "ff": "", "of": "", "et": "",
}
AREAS = ["Presupuesto", "Patrimonio", "Tesorería"]
IMPUTACIONES = ["CON IMPUTACIÓN PRESUPUESTARIA", "SIN IMPUTACIÓN PRESUPUESTARIA"]
OPERACIONES = ["Balance Inicial", "Traspaso Caja-Banco", "Ajuste Contable"]
MOMENTOS_RECURSOS = ["ESTIMACIÓN", "DEVENGADO", "PERCIBIDO"]
MOMENTOS_GASTOS = ["APROPIACIÓN", "PREVENTIVO", "COMPROMISO", "DEVENGADO", "PAGADO"]


def _glosa(meta: dict) -> str:
    return meta.get("glosa") or (
        f"Registro de la planilla de sueldos y salarios correspondiente a {meta['periodo']}.")


def _monto(v: Decimal | None) -> str:
    return "" if v is None else _fmt_monto(v)


def _doc(buf, apaisado: bool) -> SimpleDocTemplate:
    tam = landscape(letter) if apaisado else letter
    m = 22 if apaisado else 36
    return SimpleDocTemplate(buf, pagesize=tam, leftMargin=m, rightMargin=m,
                             topMargin=m, bottomMargin=m)


# --------------------------------------------------------------------------- #
# Empresa privada
# --------------------------------------------------------------------------- #
_SANGRIA = {0: 0, 1: 14, 2: 22, 3: 36}


def pdf_asiento_privado(asiento: dict, meta: dict) -> bytes:
    """Asiento contable de la planilla (empresa privada), carta vertical."""
    buf = io.BytesIO()
    doc = _doc(buf, apaisado=False)
    ancho = letter[0] - 72
    st_t = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=12, leading=15,
                          textColor=colors.HexColor("#" + COLOR_HEADER))
    st_n = ParagraphStyle("n", fontName="Helvetica", fontSize=8.5, leading=11)
    st_b = ParagraphStyle("b", fontName="Helvetica-Bold", fontSize=10, leading=13)

    historia: list = [Paragraph(meta["empresa"], st_t)]
    if meta.get("nit"):
        historia.append(Paragraph(f"NIT: {meta['nit']}", st_n))
    historia += [
        Paragraph("ASIENTO CONTABLE - REGISTRO DE LA PLANILLA DE SUELDOS Y SALARIOS", st_b),
        Paragraph(f"Periodo: {meta['periodo']}  |  Fecha: "
                  f"{_fmt_fecha(meta.get('fecha_asiento') or meta['fecha_corte'])}"
                  f"  |  Asiento N°: {meta.get('numero_asiento', '')}", st_n),
        Spacer(1, 10)]

    datos: list[list[Any]] = [["CUENTAS", "PARCIAL", "DEBE", "HABER"]]
    estilos: list[tuple] = []
    for l in asiento["lineas"]:
        txt = l["cuenta"]
        if l["subrayar"]:
            txt = f"<u>{txt}</u>"
        sty = ParagraphStyle("c", fontName="Helvetica-Bold" if l["nivel"] == 0 else "Helvetica",
                             fontSize=9, leading=11, leftIndent=_SANGRIA[l["nivel"]])
        datos.append([Paragraph(txt, sty), _monto(l["parcial"]), _monto(l["debe"]),
                      _monto(l["haber"])])
    datos.append(["TOTAL", "", _monto(asiento["total_debe"]), _monto(asiento["total_haber"])])

    t = Table(datos, colWidths=[ancho - 270, 90, 90, 90])
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.8, NEGRO), ("INNERGRID", (0, 0), (-1, 0), 0.8, NEGRO),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, NEGRO),
        ("LINEBEFORE", (1, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
        ("FONT", (1, 1), (-1, -1), "Helvetica", 9),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("BACKGROUND", (0, -1), (-1, -1), GRIS),
        ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 9),
        ("LINEABOVE", (0, -1), (-1, -1), 0.8, NEGRO),
    ] + estilos))
    historia.append(t)
    if not asiento["cuadra"]:
        historia.append(Paragraph("<b>ATENCIÓN: el asiento no cuadra.</b>", st_n))
    historia += [Spacer(1, 8), Paragraph(f"<b>Glosa:</b> {_glosa(meta)}", st_n),
                 Spacer(1, 40), KeepTogether(_firmas_pdf(ancho))]
    doc.build(historia, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


def _hoja_encabezado(ws, meta: dict, titulo: str) -> None:
    ws["A1"] = meta["empresa"]
    ws["A1"].font = Font(bold=True, size=14, color=COLOR_HEADER)
    ws["A2"] = f"NIT: {meta['nit']}" if meta.get("nit") else ""
    ws["A3"] = titulo
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = (f"Periodo: {meta['periodo']}  |  Fecha: "
                f"{_fmt_fecha(meta.get('fecha_asiento') or meta['fecha_corte'])}")


def excel_asiento_privado(asiento: dict, meta: dict) -> bytes:
    """Mismo asiento en Excel (montos como números, sangría por nivel)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Asiento"
    _hoja_encabezado(ws, meta, "ASIENTO CONTABLE - REGISTRO DE LA PLANILLA DE SUELDOS Y SALARIOS")
    borde = Side(style="thin", color="888888")
    for j, h in enumerate(["CUENTAS", "PARCIAL", "DEBE", "HABER"], start=1):
        c = ws.cell(6, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLOR_HEADER)
        c.alignment = Alignment(horizontal="center")
    r = 7
    for l in asiento["lineas"]:
        c = ws.cell(r, 1, l["cuenta"])
        c.alignment = Alignment(indent=_SANGRIA[l["nivel"]] // 6)
        c.font = Font(bold=l["nivel"] == 0, underline="single" if l["subrayar"] else None)
        for j, k in ((2, "parcial"), (3, "debe"), (4, "haber")):
            if l[k] is not None:
                cc = ws.cell(r, j, float(l[k]))
                cc.number_format = "#,##0.00"
        r += 1
    ws.cell(r, 1, "TOTAL")
    ws.cell(r, 3, f"=SUM(C7:C{r - 1})").number_format = "#,##0.00"
    ws.cell(r, 4, f"=SUM(D7:D{r - 1})").number_format = "#,##0.00"
    for j in range(1, 5):
        ws.cell(r, j).font = Font(bold=True)
        ws.cell(r, j).fill = PatternFill("solid", fgColor=COLOR_TOTAL)
        ws.cell(r, j).border = Border(top=borde)
    ws.cell(r + 2, 1, f"Glosa: {_glosa(meta)}")
    for col, w in zip("ABCD", (52, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Sector público: comprobante patrimonial
# --------------------------------------------------------------------------- #
_P_PEQ = ParagraphStyle("pq", fontName="Helvetica", fontSize=6.5, leading=7.6)
_P_NEG = ParagraphStyle("pn", fontName="Helvetica-Bold", fontSize=7, leading=8.4)


def _opciones(opciones: Sequence[str], marcadas: Sequence[str], ancho_etq: float) -> Table:
    """Lista de opciones con casilla; la marcada lleva una X."""
    marc = {m.upper() for m in marcadas if m}
    filas = [[Paragraph(o, ParagraphStyle("o", parent=_P_PEQ, alignment=2)),
              "X" if o.upper() in marc else ""] for o in opciones]
    t = Table(filas, colWidths=[ancho_etq, 24], rowHeights=[11] * len(filas))
    t.setStyle(TableStyle([
        ("BOX", (1, 0), (1, -1), 0.8, NEGRO), ("INNERGRID", (1, 0), (1, -1), 0.5, NEGRO),
        ("FONT", (1, 0), (1, -1), "Helvetica-Bold", 7), ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return t


def pdf_comprobante_publico(asiento: dict, meta: dict) -> bytes:
    """Comprobante patrimonial (sector público), carta apaisado."""
    cmp_ = dict(COMPROBANTE_POR_DEFECTO)
    cmp_.update(meta.get("comprobante") or {})
    buf = io.BytesIO()
    doc = _doc(buf, apaisado=True)
    W = landscape(letter)[0] - 44
    fecha = _fmt_fecha(meta.get("fecha_asiento") or meta["fecha_corte"])

    def caja(datos, **kw):
        t = Table(datos, **kw)
        return t

    # 1. Título
    titulo = Table([[Paragraph(f"<b>{meta['empresa'].upper()}</b>",
                               ParagraphStyle("x", fontName="Helvetica-Bold", fontSize=11,
                                              alignment=1))],
                    [Paragraph("<b>COMPROBANTE PATRIMONIAL</b>",
                               ParagraphStyle("y", fontName="Helvetica-Bold", fontSize=10,
                                              alignment=1))]], colWidths=[W])
    titulo.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO)]))

    # 2. Tipo / imputación / fecha
    c1 = [Paragraph(f"<b>Comprobante N°:</b> {cmp_['numero']}", _P_PEQ),
          _opciones(AREAS, cmp_["areas"], 80)]
    c2 = [Spacer(1, 12), _opciones(IMPUTACIONES, [cmp_["imputacion"]], 170)]
    c3 = [Paragraph(f"<b>Fecha:</b> {fecha}", _P_PEQ),
          _opciones(OPERACIONES, [cmp_["operacion"] or ""], 100)]
    info = Table([[c1, c2, c3]], colWidths=[W * 0.3, W * 0.4, W * 0.3])
    info.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO),
                              ("VALIGN", (0, 0), (-1, -1), "TOP")]))

    # 3. Momentos de registro
    rec = [Paragraph("<b>MOMENTOS DE REGISTRO DE RECURSOS</b>", _P_NEG),
           _opciones(MOMENTOS_RECURSOS, [cmp_["momento_recursos"] or ""], 90)]
    gas = [Paragraph("<b>MOMENTOS DE REGISTRO DE GASTOS</b>", _P_NEG),
           _opciones(MOMENTOS_GASTOS, [cmp_["momento_gastos"] or ""], 90)]
    momentos = Table([[rec, gas]], colWidths=[W / 2, W / 2])
    momentos.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO),
                                  ("LINEAFTER", (0, 0), (0, 0), 0.8, NEGRO),
                                  ("VALIGN", (0, 0), (-1, -1), "TOP")]))

    # 4. Beneficiario y respaldo
    benef = Table([[Paragraph(f"<b>BENEFICIARIO:</b> {cmp_['beneficiario']}", _P_PEQ)],
                   [Paragraph(f"<b>DOCUMENTO DE RESPALDO:</b> {cmp_['respaldo']}", _P_PEQ)]],
                  colWidths=[W])
    benef.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO)]))

    # 5. Detalle
    anchos = [72, 46, 22, 22, 22, 140, 52, 176, 62, 62, 62]
    anchos[5] += W - sum(anchos)
    enc = ["APERTURA\nPROGRAMÁTICA", "RUBRO /\nPARTIDA", "FF", "OF", "ET",
           "DESCRIPCIÓN PARTIDA PRESUPUESTARIA", "CÓDIGO\nCONTABLE",
           "DESCRIPCIÓN CUENTA CONTABLE", "PARCIAL", "DEBE", "HABER"]
    filas: list[list[Any]] = [[Paragraph(e.replace("\n", "<br/>"),
                                         ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=6.5,
                                                        leading=7.5, alignment=1)) for e in enc]]
    primera = True
    for l in asiento["lineas"]:
        con_ppto = bool(l["rubro"])
        filas.append([
            cmp_["apertura"] if (con_ppto and primera) else "",
            l["rubro"], cmp_["ff"] if con_ppto else "", cmp_["of"] if con_ppto else "",
            cmp_["et"] if con_ppto else "", l["partida"], l["codigo"], l["cuenta"],
            _monto(l["parcial"]), _monto(l["debe"]), _monto(l["haber"])])
        if con_ppto:
            primera = False
    n_det = len(filas) - 1
    filas.append(["", "", "", "", "", "", "", "TOTAL", "", _monto(asiento["total_debe"]),
                  _monto(asiento["total_haber"])])
    det = Table(filas, colWidths=anchos, repeatRows=1)
    det.setStyle(TableStyle([
        ("FONT", (0, 1), (-1, -1), "Helvetica", 6.5),
        ("BOX", (0, 0), (-1, -1), 0.8, NEGRO), ("GRID", (0, 0), (-1, 0), 0.8, NEGRO),
        ("LINEBEFORE", (1, 1), (-1, -2), 0.5, NEGRO),
        ("ALIGN", (0, 1), (1, -1), "CENTER"), ("ALIGN", (2, 1), (4, -1), "CENTER"),
        ("ALIGN", (6, 1), (6, -1), "CENTER"), ("ALIGN", (8, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 1), (-1, -1), 1.1), ("BOTTOMPADDING", (0, 1), (-1, -1), 1.1),
        ("BOX", (0, -1), (-1, -1), 0.8, NEGRO),
        ("FONT", (7, -1), (-1, -1), "Helvetica-Bold", 7), ("ALIGN", (7, -1), (7, -1), "RIGHT"),
        ("BOX", (9, -1), (9, -1), 0.8, NEGRO), ("BOX", (10, -1), (10, -1), 0.8, NEGRO),
    ]))
    glosa = Table([[Paragraph(f"<b>GLOSA:</b> {_glosa(meta)}", _P_PEQ)]], colWidths=[W],
                  rowHeights=[26])
    glosa.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO), ("VALIGN", (0, 0), (-1, -1), "TOP")]))

    def firma(etq, cargo):
        t = Table([[Paragraph(f"<b>{etq}</b>", _P_PEQ)], [""], [Paragraph(f"<b>{cargo}</b>",
                  ParagraphStyle("c", fontName="Helvetica-Bold", fontSize=7, alignment=1))]],
                  colWidths=[W * 0.36], rowHeights=[14, 26, 16])
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, NEGRO),
                               ("LINEABOVE", (0, 2), (-1, 2), 0.8, NEGRO)]))
        return t

    firmas = Table([[firma("ELABORADO POR:", "CONTADOR"), "",
                     firma("REVISADO POR:", "GERENTE FINANCIERO")]],
                   colWidths=[W * 0.36, W * 0.12, W * 0.36])
    firmas.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER")]))

    historia = [titulo, info, momentos, benef, Spacer(1, 6), det, glosa, Spacer(1, 8),
                KeepTogether(firmas)]
    if not asiento["cuadra"]:
        historia.append(Paragraph("<b>ATENCIÓN: el comprobante no cuadra.</b>", _P_PEQ))
    doc.build(historia, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


def excel_comprobante_publico(asiento: dict, meta: dict) -> bytes:
    """Comprobante patrimonial en Excel."""
    cmp_ = dict(COMPROBANTE_POR_DEFECTO)
    cmp_.update(meta.get("comprobante") or {})
    wb = Workbook()
    ws = wb.active
    ws.title = "Comprobante"
    _hoja_encabezado(ws, meta, "COMPROBANTE PATRIMONIAL")
    ws["A2"] = ""
    ws["A5"] = (f"Comprobante N°: {cmp_['numero']}  |  Áreas: {', '.join(cmp_['areas'])}  |  "
                f"{cmp_['imputacion']}")
    ws["A6"] = (f"Momento de registro de gastos: {cmp_['momento_gastos'] or '-'}  |  "
                f"Momento de registro de recursos: {cmp_['momento_recursos'] or '-'}  |  "
                f"Operación: {cmp_['operacion'] or '-'}")
    ws["A7"] = f"Beneficiario: {cmp_['beneficiario']}"
    ws["A8"] = f"Documento de respaldo: {cmp_['respaldo']}"
    enc = ["Apertura programática", "Rubro/Partida", "FF", "OF", "ET", "Partida presupuestaria",
           "Código contable", "Descripción cuenta contable", "PARCIAL", "DEBE", "HABER"]
    for j, h in enumerate(enc, start=1):
        c = ws.cell(10, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COLOR_HEADER)
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[10].height = 30
    r, primera = 11, True
    for l in asiento["lineas"]:
        con = bool(l["rubro"])
        vals = [cmp_["apertura"] if (con and primera) else "", l["rubro"] or None,
                cmp_["ff"] if con else "", cmp_["of"] if con else "", cmp_["et"] if con else "",
                l["partida"], l["codigo"], l["cuenta"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j, v)
        for j, k in ((9, "parcial"), (10, "debe"), (11, "haber")):
            if l[k] is not None:
                ws.cell(r, j, float(l[k])).number_format = "#,##0.00"
        primera = primera and not con
        r += 1
    ws.cell(r, 8, "TOTAL").font = Font(bold=True)
    ws.cell(r, 10, f"=SUM(J11:J{r - 1})").number_format = "#,##0.00"
    ws.cell(r, 11, f"=SUM(K11:K{r - 1})").number_format = "#,##0.00"
    for j in range(8, 12):
        ws.cell(r, j).font = Font(bold=True)
        ws.cell(r, j).fill = PatternFill("solid", fgColor=COLOR_TOTAL)
    ws.cell(r + 2, 1, f"GLOSA: {_glosa(meta)}")
    ws.cell(r + 5, 2, "ELABORADO POR: CONTADOR")
    ws.cell(r + 5, 8, "REVISADO POR: GERENTE FINANCIERO")
    for col, w in zip("ABCDEFGHIJK", (16, 12, 6, 6, 6, 40, 12, 44, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
