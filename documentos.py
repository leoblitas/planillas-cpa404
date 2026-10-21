"""
documentos.py
=============
Generación de las dos planillas en Excel (.xlsx) y PDF (tamaño carta, apaisado):

* Planilla Laboral (sueldos y salarios)  + hoja/tabla de aportes patronales
* Planilla Impositiva (RC-IVA)

Todas las funciones públicas devuelven ``bytes`` listos para
``st.download_button`` o para guardarse en disco.

``filas`` es la lista de dicts que produce ``planilla.calcular_todos``.
``meta`` es un dict con: ``empresa``, ``nit``, ``periodo`` (texto, p. ej.
"Octubre 2026"), ``fecha_corte`` (date), ``parametros`` (dict de parámetros),
``ufv_anterior`` y ``ufv_actual``.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

# Tipos de columna: n = correlativo, txt = texto, date = fecha, int = entero,
# m = monto en Bs (se totaliza), sig = espacio para firma.
Col = tuple  # (encabezado, clave, ancho_pdf_pt, tipo)

COLOR_HEADER = "1F3A5F"
COLOR_GROUP = "2E5C8A"
COLOR_TOTAL = "E8EEF5"
COLOR_ZEBRA = "F6F8FB"


# --------------------------------------------------------------------------- #
# Definición de columnas
# --------------------------------------------------------------------------- #
def _cols_laboral_xlsx() -> tuple[list[tuple[str, int]], list[Col]]:
    grupos = [("TRABAJADOR", 4), ("INGRESOS", 8), ("DESCUENTOS", 7), ("PAGO", 2)]
    cols = [
        ("N°", "n", 5, "n"), ("CI", "ci", 12, "txt"),
        ("Apellidos y nombres", "nombre", 38, "txt"),
        ("Fecha de ingreso", "fecha_ingreso", 12, "date"),
        ("Haber básico", "haber_basico", 13, "m"),
        ("Años antig.", "anios_antiguedad", 8, "int"),
        ("Bono antigüedad", "bono_antiguedad", 13, "m"),
        ("Horas extra", "pago_horas_extra", 12, "m"),
        ("Rec. nocturno", "pago_nocturno", 12, "m"),
        ("Domingos", "pago_domingos", 12, "m"),
        ("Feriados", "pago_feriados", 12, "m"),
        ("TOTAL GANADO", "total_ganado", 14, "m"),
        ("Cuenta personal", "aporte_cuenta_personal", 13, "m"),
        ("Riesgo común", "aporte_riesgo_comun", 12, "m"),
        ("Comisión", "aporte_comision", 11, "m"),
        ("Solidario", "aporte_solidario_laboral", 11, "m"),
        ("ANS", "aporte_nacional_solidario", 11, "m"),
        ("RC-IVA", "rciva_retener", 12, "m"),
        ("TOTAL DESCUENTOS", "total_descuentos", 14, "m"),
        ("LÍQUIDO PAGABLE", "liquido_pagable", 14, "m"),
        ("Firma", None, 24, "sig"),
    ]
    return grupos, cols


def _cols_laboral_pdf() -> tuple[list[tuple[str, int]], list[Col]]:
    grupos = [("TRABAJADOR", 3), ("INGRESOS", 8), ("DESCUENTOS", 4), ("PAGO", 2)]
    cols = [
        ("N°", "n", 16, "n"), ("CI", "ci", 46, "txt"),
        ("Apellidos y nombres", "nombre", 86, "txt"),
        ("Haber básico", "haber_basico", 42, "m"),
        ("Años antig.", "anios_antiguedad", 22, "int"),
        ("Bono antig.", "bono_antiguedad", 38, "m"),
        ("Horas extra", "pago_horas_extra", 38, "m"),
        ("Rec. noct.", "pago_nocturno", 34, "m"),
        ("Domingos", "pago_domingos", 36, "m"),
        ("Feriados", "pago_feriados", 36, "m"),
        ("TOTAL GANADO", "total_ganado", 44, "m"),
        ("Aportes AFP", "aportes_afp", 40, "m"),
        ("ANS", "aporte_nacional_solidario", 34, "m"),
        ("RC-IVA", "rciva_retener", 38, "m"),
        ("TOTAL DESC.", "total_descuentos", 42, "m"),
        ("LÍQUIDO PAGABLE", "liquido_pagable", 46, "m"),
        ("Firma", None, 54, "sig"),
    ]
    return grupos, cols


def _cols_patronal(ancho_pdf: bool) -> tuple[list[tuple[str, int]], list[Col]]:
    grupos = [("TRABAJADOR", 3), ("APORTES PATRONALES", 5), ("PREVISIONES", 2)]
    w = (lambda p, x: p) if ancho_pdf else (lambda p, x: x)
    cols = [
        ("N°", "n", w(24, 5), "n"),
        ("Apellidos y nombres", "nombre", w(150, 38), "txt"),
        ("Total ganado", "total_ganado", w(70, 14), "m"),
        ("CNS", "patronal_cns", w(66, 13), "m"),
        ("Solidario patronal", "patronal_solidario", w(66, 13), "m"),
        ("Riesgo profesional", "patronal_riesgo_profesional", w(66, 13), "m"),
        ("Pro-vivienda", "patronal_provivienda", w(66, 13), "m"),
        ("TOTAL PATRONAL", "aportes_patronales", w(74, 14), "m"),
        ("Aguinaldo", "aguinaldo", w(66, 13), "m"),
        ("Indemnización", "indemnizacion", w(66, 13), "m"),
    ]
    return grupos, cols


def _cols_impositiva(pdf: bool, p: dict) -> tuple[list[tuple[str, int]], list[Col]]:
    tasa = f"{Decimal(str(p['T_RCIVA'])) * 100:.0f}%"
    smn_ex = f"{float(p['SMN_EXENTO']):g}"
    grupos = [("TRABAJADOR", 3), ("BASE IMPONIBLE", 5),
              ("IMPUESTO Y CRÉDITOS", 6), ("RESULTADO", 2)]
    if pdf:
        w = lambda a, b: a
        nums = 44
    else:
        w = lambda a, b: b
        nums = 15
    cols = [
        ("N°", "n", w(16, 5), "n"), ("CI", "ci", w(48, 12), "txt"),
        ("Apellidos y nombres", "nombre", w(100, 38), "txt"),
        ("Total ganado", "total_ganado", w(nums, nums), "m"),
        ("Aportes laborales", "aportes_laborales", w(nums, nums), "m"),
        ("Base imponible", "base_imponible", w(nums, nums), "m"),
        (f"Mínimo no imp. ({smn_ex} SMN)", "minimo_no_imponible", w(nums, nums), "m"),
        ("Exceso sujeto a impuesto", "exceso_imponible", w(nums, nums), "m"),
        (f"Impuesto determinado ({tasa})", "impuesto_determinado", w(nums, nums), "m"),
        ("Crédito dependiente", "credito_dependiente", w(nums, nums), "m"),
        ("Facturas F-110 (monto)", "facturas_f110", w(nums, nums), "m"),
        ("Crédito F-110", "credito_f110", w(nums, nums), "m"),
        ("Saldo anterior actualiz.", "saldo_favor_anterior_actualizado", w(nums, nums), "m"),
        ("RC-IVA neto", "rciva_neto", w(nums, nums), "m"),
        ("RC-IVA A RETENER", "rciva_retener", w(nums, nums), "m"),
        ("SALDO A FAVOR siguiente", "saldo_favor_nuevo", w(nums, nums), "m"),
    ]
    return grupos, cols


# --------------------------------------------------------------------------- #
# Utilidades de formato
# --------------------------------------------------------------------------- #
def _fmt_monto(v: Any) -> str:
    """1234.5 -> '1.234,50' (formato boliviano)."""
    s = f"{Decimal(str(v)):,.2f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def _fmt_fecha(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _totales(filas: Sequence[dict], cols: Sequence[Col]) -> dict[str, Decimal]:
    return {k: sum((Decimal(str(f[k])) for f in filas), Decimal("0"))
            for _, k, _, t in cols if t == "m"}


def _notas_impositiva(meta: dict) -> list[str]:
    p = meta["parametros"]
    smn = Decimal(str(p["SMN"]))
    tasa = Decimal(str(p["T_RCIVA"]))
    cred = tasa * Decimal(str(p["SMN_CREDITO_DEP"])) * smn
    factor = max(Decimal(str(meta["ufv_actual"])) / Decimal(str(meta["ufv_anterior"])),
                 Decimal("1"))
    return [
        f"Salario Mínimo Nacional (SMN): Bs {_fmt_monto(smn)}. "
        f"Mínimo no imponible: {float(p['SMN_EXENTO']):g} SMN = "
        f"Bs {_fmt_monto(Decimal(str(p['SMN_EXENTO'])) * smn)}. "
        f"Tasa RC-IVA: {tasa * 100:.0f}%.",
        f"Crédito fiscal por dependiente: {tasa * 100:.0f}% de "
        f"{float(p['SMN_CREDITO_DEP']):g} SMN = Bs {_fmt_monto(cred)}. "
        f"Crédito F-110: {tasa * 100:.0f}% del monto de facturas presentado.",
        f"Actualización del saldo a favor anterior por UFV: "
        f"{meta['ufv_actual']} / {meta['ufv_anterior']} = factor {factor:.6f}.",
    ]


def _encabezado_textos(meta: dict, titulo: str) -> list[str]:
    return [meta["empresa"], f"NIT: {meta['nit']}" if meta.get("nit") else "",
            titulo,
            f"Periodo: {meta['periodo']}  |  Fecha de corte: "
            f"{_fmt_fecha(meta['fecha_corte'])}  |  Expresado en bolivianos (Bs)"]


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _hoja_excel(ws, meta: dict, titulo: str, grupos, cols, filas,
                notas: Sequence[str] = (), firmas: bool = True) -> None:
    borde = Side(style="thin", color="B8C2CC")
    caja = Border(left=borde, right=borde, top=borde, bottom=borde)
    ncols = len(cols)

    empresa, nit, tit, periodo = _encabezado_textos(meta, titulo)
    ws["A1"], ws["A2"], ws["A3"], ws["A4"] = empresa, nit, tit, periodo
    ws["A1"].font = Font(bold=True, size=14, color=COLOR_HEADER)
    ws["A3"].font = Font(bold=True, size=12)
    ws["A2"].font = ws["A4"].font = Font(size=10, color="444444")

    fila_g, fila_h, fila_ini = 6, 7, 8

    # Grupos
    col = 1
    for nombre, span in grupos:
        ws.cell(fila_g, col, nombre)
        if span > 1:
            ws.merge_cells(start_row=fila_g, start_column=col,
                           end_row=fila_g, end_column=col + span - 1)
        for k in range(col, col + span):
            c = ws.cell(fila_g, k)
            c.fill = PatternFill("solid", fgColor=COLOR_GROUP)
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = caja
        col += span

    # Encabezados
    for j, (titulo_col, _, ancho, _t) in enumerate(cols, start=1):
        c = ws.cell(fila_h, j, titulo_col)
        c.fill = PatternFill("solid", fgColor=COLOR_HEADER)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = caja
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila_h].height = 42

    # Datos
    for i, f in enumerate(filas):
        r = fila_ini + i
        for j, (_, clave, _, tipo) in enumerate(cols, start=1):
            c = ws.cell(r, j)
            c.border = caja
            if tipo == "sig":
                continue
            v = f[clave]
            if tipo == "m":
                c.value = float(v)
                c.number_format = "#,##0.00"
            elif tipo == "date":
                c.value = v
                c.number_format = "DD/MM/YYYY"
                c.alignment = Alignment(horizontal="center")
            elif tipo in ("int", "n"):
                c.value = int(v)
                c.alignment = Alignment(horizontal="center")
            else:
                c.value = v
        ws.row_dimensions[r].height = 22 if any(t == "sig" for *_, t in cols) else 16

    # Totales con fórmulas SUM
    fin = fila_ini + len(filas) - 1
    r_tot = fin + 1
    primera_txt = next(j for j, (*_, t) in enumerate(cols, start=1) if t == "txt")
    for j, (_, clave, _, tipo) in enumerate(cols, start=1):
        c = ws.cell(r_tot, j)
        c.fill = PatternFill("solid", fgColor=COLOR_TOTAL)
        c.font = Font(bold=True)
        c.border = Border(left=borde, right=borde, top=Side(style="medium"), bottom=borde)
        if tipo == "m" and filas:
            letra = get_column_letter(j)
            c.value = f"=SUM({letra}{fila_ini}:{letra}{fin})"
            c.number_format = "#,##0.00"
    ws.cell(r_tot, primera_txt + 1 if cols[primera_txt][3] == "txt" and
            cols[primera_txt][1] == "ci" else primera_txt, "TOTALES")

    # Notas y firmas
    r = r_tot + 2
    for n in notas:
        ws.cell(r, 1, n).font = Font(italic=True, size=9, color="444444")
        r += 1
    if firmas:
        r += 3
        for k, etiqueta in enumerate(["Elaborado por", "Revisado por", "Aprobado por"]):
            colf = 3 + k * 4 if ncols >= 14 else 2 + k * 3
            c = ws.cell(r, colf, etiqueta)
            c.alignment = Alignment(horizontal="center")
            c.border = Border(top=Side(style="thin"))
            c.font = Font(size=9)

    # Impresión: carta apaisada, ajustar al ancho
    ws.freeze_panes = ws.cell(fila_ini, 4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{fila_g}:{fila_h}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.sheet_view.showGridLines = False


def _a_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_planilla_laboral(filas: Sequence[dict], meta: dict) -> bytes:
    """Libro con la Planilla Laboral y una segunda hoja de aportes patronales."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilla Laboral"
    g, c = _cols_laboral_xlsx()
    _hoja_excel(ws, meta, "PLANILLA DE SUELDOS Y SALARIOS", g, c, filas)
    ws2 = wb.create_sheet("Aportes patronales")
    g2, c2 = _cols_patronal(ancho_pdf=False)
    _hoja_excel(ws2, meta, "APORTES PATRONALES Y PREVISIONES", g2, c2, filas, firmas=False)
    return _a_bytes(wb)


def excel_planilla_impositiva(filas: Sequence[dict], meta: dict) -> bytes:
    """Libro con la Planilla Impositiva (RC-IVA)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilla Impositiva"
    g, c = _cols_impositiva(False, meta["parametros"])
    _hoja_excel(ws, meta, "PLANILLA IMPOSITIVA - RC-IVA (DEPENDIENTES)", g, c, filas,
                notas=_notas_impositiva(meta))
    return _a_bytes(wb)


# --------------------------------------------------------------------------- #
# PDF (carta apaisada)
# --------------------------------------------------------------------------- #
_ESTILO_ENC = ParagraphStyle("enc", fontName="Helvetica-Bold", fontSize=6.3, leading=7.2,
                             alignment=1, textColor=colors.white)
_ESTILO_NOMBRE = ParagraphStyle("nom", fontName="Helvetica", fontSize=6.5, leading=7.5)
_ESTILO_NOTA = ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=7, leading=9,
                              textColor=colors.HexColor("#444444"))


def _tabla_pdf(grupos, cols, filas) -> Table:
    ancho_total = sum(c[2] for c in cols)
    anchos = [c[2] for c in cols]
    # fila 0: grupos, fila 1: encabezados
    fila_grupos: list[Any] = []
    for nombre, span in grupos:
        fila_grupos.append(Paragraph(nombre, _ESTILO_ENC))
        fila_grupos.extend([""] * (span - 1))
    fila_enc = [Paragraph(c[0], _ESTILO_ENC) for c in cols]

    tot = _totales(filas, cols)
    datos: list[list[Any]] = [fila_grupos, fila_enc]
    for f in filas:
        fila = []
        for _, clave, _, tipo in cols:
            if tipo == "sig":
                fila.append("")
            elif tipo == "m":
                fila.append(_fmt_monto(f[clave]))
            elif tipo == "txt" and clave == "nombre":
                fila.append(Paragraph(str(f[clave]), _ESTILO_NOMBRE))
            elif tipo == "date":
                fila.append(_fmt_fecha(f[clave]))
            else:
                fila.append(str(f[clave]))
        datos.append(fila)
    fila_tot = []
    pos_label = next(j for j, c in enumerate(cols) if c[3] == "txt" and c[1] == "nombre")
    for j, (_, clave, _, tipo) in enumerate(cols):
        if tipo == "m":
            fila_tot.append(_fmt_monto(tot[clave]))
        elif j == pos_label:
            fila_tot.append("TOTALES")
        else:
            fila_tot.append("")
    datos.append(fila_tot)

    t = Table(datos, colWidths=anchos, repeatRows=2)
    estilo = [
        ("FONT", (0, 2), (-1, -1), "Helvetica", 6.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + COLOR_GROUP)),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#" + COLOR_HEADER)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C2CC")),
        ("ALIGN", (0, 2), (-1, -1), "RIGHT"),
        ("TOPPADDING", (0, 1), (-1, -1), 3), ("BOTTOMPADDING", (0, 1), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("ROWBACKGROUNDS", (0, 2), (-1, -2), [colors.white, colors.HexColor("#" + COLOR_ZEBRA)]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#" + COLOR_TOTAL)),
        ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 6.5),
        ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
    ]
    col = 0
    for _, span in grupos:
        if span > 1:
            estilo.append(("SPAN", (col, 0), (col + span - 1, 0)))
        col += span
    for j, c in enumerate(cols):          # texto y fecha centrados/izquierda
        if c[3] in ("n", "int", "date"):
            estilo.append(("ALIGN", (j, 2), (j, -2), "CENTER"))
        if c[3] == "txt":
            estilo.append(("ALIGN", (j, 2), (j, -2), "LEFT"))
    estilo.append(("ALIGN", (0, 0), (-1, 1), "CENTER"))
    if ancho_total > 0:
        t.setStyle(TableStyle(estilo))
    return t


def _firmas_pdf(ancho: float) -> Table:
    """Tres líneas de firma separadas entre sí."""
    seg, hueco = ancho * 0.24, ancho * 0.14
    t = Table([["Elaborado por", "", "Revisado por", "", "Aprobado por"]],
              colWidths=[seg, hueco, seg, hueco, seg], rowHeights=[16])
    t.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 7.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("LINEABOVE", (0, 0), (0, 0), 0.6, colors.black),
        ("LINEABOVE", (2, 0), (2, 0), 0.6, colors.black),
        ("LINEABOVE", (4, 0), (4, 0), 0.6, colors.black),
    ]))
    return t


def _encabezado_pdf(meta: dict, titulo: str) -> list:
    empresa, nit, tit, periodo = _encabezado_textos(meta, titulo)
    e1 = ParagraphStyle("e1", fontName="Helvetica-Bold", fontSize=12, leading=14,
                        textColor=colors.HexColor("#" + COLOR_HEADER))
    e2 = ParagraphStyle("e2", fontName="Helvetica", fontSize=8, leading=10)
    e3 = ParagraphStyle("e3", fontName="Helvetica-Bold", fontSize=10, leading=13)
    out = [Paragraph(empresa, e1)]
    if nit:
        out.append(Paragraph(nit, e2))
    out += [Paragraph(tit, e3), Paragraph(periodo, e2), Spacer(1, 8)]
    return out


def _pie(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 16, f"Página {doc.page}")
    canvas.restoreState()


def _doc_pdf(buf) -> SimpleDocTemplate:
    return SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=28,
                             rightMargin=28, topMargin=28, bottomMargin=28)


def pdf_planilla_laboral(filas: Sequence[dict], meta: dict) -> bytes:
    """PDF carta apaisado: Planilla Laboral y, en página aparte, aportes patronales."""
    buf = io.BytesIO()
    doc = _doc_pdf(buf)
    g, c = _cols_laboral_pdf()
    historia = _encabezado_pdf(meta, "PLANILLA DE SUELDOS Y SALARIOS")
    historia.append(_tabla_pdf(g, c, filas))
    historia += [Spacer(1, 26), KeepTogether(_firmas_pdf(sum(x[2] for x in c)))]
    g2, c2 = _cols_patronal(ancho_pdf=True)
    historia += [PageBreak()] + _encabezado_pdf(meta, "APORTES PATRONALES Y PREVISIONES")
    historia.append(_tabla_pdf(g2, c2, filas))
    doc.build(historia, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


def pdf_planilla_impositiva(filas: Sequence[dict], meta: dict) -> bytes:
    """PDF carta apaisado de la Planilla Impositiva (RC-IVA)."""
    buf = io.BytesIO()
    doc = _doc_pdf(buf)
    g, c = _cols_impositiva(True, meta["parametros"])
    historia = _encabezado_pdf(meta, "PLANILLA IMPOSITIVA - RC-IVA (DEPENDIENTES)")
    historia.append(_tabla_pdf(g, c, filas))
    historia.append(Spacer(1, 8))
    for nota in _notas_impositiva(meta):
        historia.append(Paragraph(nota, _ESTILO_NOTA))
    historia += [Spacer(1, 26), KeepTogether(_firmas_pdf(sum(x[2] for x in c)))]
    doc.build(historia, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()
