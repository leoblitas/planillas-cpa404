"""
planilla.py
===========
Capa intermedia entre la interfaz (Streamlit) y el motor de cálculo
(``calculos.py``). No depende de Streamlit.

* ``cargar_datos_excel``: lee parámetros, tabla de antigüedad y empleados del
  libro de Excel (hojas ``Parametros_LEY``, ``Antiguedad`` y ``Empleados``).
* ``calcular_todos``: calcula la planilla de una lista de empleados y devuelve
  filas planas listas para los documentos, más la lista de errores por
  empleado (un dato malo no detiene al resto).
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from typing import Any, BinaryIO, Union

from openpyxl import load_workbook

import calculos as c

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

# Variables del mes que se ingresan por empleado
COLUMNAS_VARIABLES = ["Horas_Extra", "Horas_Nocturnas", "Dias_Domingo",
                      "Dias_Feriado", "Facturas_F110", "Saldo_Favor_Anterior"]


# --------------------------------------------------------------------------- #
# Lectura del Excel
# --------------------------------------------------------------------------- #
def _leer_tabla(ws) -> list[dict[str, Any]]:
    """Lee una hoja como lista de dicts. El encabezado es la primera fila con
    al menos dos celdas llenas (tolera columnas/filas vacías al inicio)."""
    filas = list(ws.iter_rows(values_only=True))
    for i, fila in enumerate(filas):
        if sum(v is not None for v in fila) >= 2:
            cols = [(j, str(v).strip()) for j, v in enumerate(fila) if v is not None]
            registros = []
            for resto in filas[i + 1:]:
                if all(v is None for v in resto):
                    continue
                registros.append({n: resto[j] if j < len(resto) else None
                                  for j, n in cols})
            return registros
    return []


def cargar_datos_excel(archivo: Union[str, BinaryIO]
                       ) -> tuple[list[dict], list[dict], list[dict]]:
    """Devuelve ``(filas_parametros, tabla_antiguedad, empleados)``.

    Las hojas se identifican por contenido/nombre, no por posición:
    parámetros = la que tiene la columna ``Codigo_Parametro``;
    antigüedad = nombre con "antig"; empleados = nombre con "emplead".
    """
    wb = load_workbook(archivo, data_only=True)
    parametros = antiguedad = empleados = None
    for ws in wb.worksheets:
        nombre = ws.title.lower()
        tabla = _leer_tabla(ws)
        if tabla and "Codigo_Parametro" in tabla[0]:
            parametros = tabla
        elif "antig" in nombre:
            antiguedad = tabla
        elif "emplead" in nombre:
            empleados = tabla
    faltan = [n for n, v in (("parámetros", parametros),
                             ("antigüedad", antiguedad),
                             ("empleados", empleados)) if v is None]
    if faltan:
        raise ValueError("El Excel no tiene la(s) hoja(s) de: " + ", ".join(faltan))
    return parametros, antiguedad, empleados


# --------------------------------------------------------------------------- #
# Cálculo masivo
# --------------------------------------------------------------------------- #
def _num(valor: Any) -> Decimal:
    """Número desde celda/tabla; vacío o NaN cuentan como 0."""
    if valor is None or valor != valor or valor == "":
        return Decimal("0")
    return c.D(valor)


def _fecha(valor: Any) -> date:
    if valor is None or valor != valor:
        raise ValueError("falta la fecha de ingreso")
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return datetime.strptime(str(valor)[:10], "%Y-%m-%d").date()


def calcular_todos(empleados: list[dict], parametros: dict, tabla_antiguedad: list[dict],
                   anio: int, mes: int, ufv_anterior: Any = 1, ufv_actual: Any = 1
                   ) -> tuple[list[dict], list[str]]:
    """Calcula la planilla de todos los empleados del periodo.

    ``empleados``: dicts con ``ID_Empleado, CI, Apellidos, Nombres,
    Fecha_Ingreso, Haber_Basico`` y las ``COLUMNAS_VARIABLES``.
    Retorna ``(filas, errores)``.
    """
    fecha_corte = c.ultimo_dia_periodo(anio, mes)
    filas: list[dict] = []
    errores: list[str] = []
    for i, e in enumerate(empleados, start=1):
        etiqueta = str(e.get("ID_Empleado") or f"fila {i}")
        try:
            haber = _num(e.get("Haber_Basico"))
            if haber <= 0:
                raise ValueError("el haber básico debe ser mayor que cero")
            res = c.calcular_planilla_empleado(
                haber_basico=haber,
                fecha_ingreso=_fecha(e.get("Fecha_Ingreso")),
                fecha_corte=fecha_corte,
                parametros=parametros,
                tabla_antiguedad=tabla_antiguedad,
                horas_extra=_num(e.get("Horas_Extra")),
                horas_nocturnas=_num(e.get("Horas_Nocturnas")),
                dias_domingo=_num(e.get("Dias_Domingo")),
                dias_feriado=_num(e.get("Dias_Feriado")),
                facturas_f110=_num(e.get("Facturas_F110")),
                saldo_favor_anterior=_num(e.get("Saldo_Favor_Anterior")),
                ufv_anterior=ufv_anterior,
                ufv_actual=ufv_actual,
            )
        except Exception as exc:  # un empleado con datos malos no detiene el resto
            errores.append(f"{etiqueta}: {exc}")
            continue
        fila = res.to_dict()
        fila.update({
            "n": len(filas) + 1,
            "id": e.get("ID_Empleado"),
            "ci": e.get("CI") or "",
            "nombre": f"{e.get('Apellidos') or ''} {e.get('Nombres') or ''}".strip(),
            "fecha_ingreso": _fecha(e.get("Fecha_Ingreso")),
            "horas_extra": _num(e.get("Horas_Extra")),
            "horas_nocturnas": _num(e.get("Horas_Nocturnas")),
            "dias_domingo": _num(e.get("Dias_Domingo")),
            "dias_feriado": _num(e.get("Dias_Feriado")),
            "facturas_f110": _num(e.get("Facturas_F110")),
            "total_descuentos": c.redondear(fila["aportes_laborales"] + fila["rciva_retener"]),
        })
        filas.append(fila)
    return filas, errores
