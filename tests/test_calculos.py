"""Pruebas del motor de cálculo y de los documentos.

Ejecutar:  python -m pytest tests -q     (o:  python tests/test_calculos.py)
Los valores esperados se verificaron a mano.
"""
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

import calculos as c  # noqa: E402
import documentos as d  # noqa: E402
import planilla as pl  # noqa: E402

EXCEL = RAIZ / "datos" / "PlanillasCPA404_BD.xlsx"
D = Decimal


def _datos():
    filas_p, tabla, empleados = pl.cargar_datos_excel(str(EXCEL))
    return c.cargar_parametros(filas_p), tabla, empleados


def test_tramos_antiguedad():
    _, tabla, _ = _datos()
    esperado = {0: 0, 1: 0, 2: 5, 4: 5, 5: 11, 7: 11, 8: 18, 10: 18, 11: 26,
                14: 26, 15: 34, 19: 34, 20: 42, 24: 42, 25: 50, 40: 50}
    for anios, pct in esperado.items():
        assert c.porcentaje_antiguedad(anios, tabla) == D(pct) / 100, anios


def test_anios_cumplidos_y_bono():
    _, tabla, _ = _datos()
    assert c.anios_cumplidos(date(2015, 4, 1), date(2026, 3, 31)) == 10
    assert c.anios_cumplidos(date(2015, 4, 1), date(2026, 4, 1)) == 11
    r = c.calcular_bono_antiguedad(date(2015, 4, 1), date(2026, 10, 31), 3300, tabla)
    assert r["anios"] == 11 and r["bono"] == D("2574.00")  # 26% de 9.900


def test_ans_por_rangos():
    par, _, _ = _datos()
    assert c.calcular_ans(30000, par) == D("482.50")   # 17.000x1,15% + 5.000x5,74%
    assert c.calcular_ans(13000, par) == D("0.00")
    # 27.000x1,15% + 15.000x5,74% + 5.000x11,48% = 310,50 + 861 + 574
    assert c.calcular_ans(40000, par) == D("1745.50")


def test_componentes_variables_emp001():
    par, _, _ = _datos()
    r = c.calcular_componentes_variables(22450, 10, 8, 1, 1, par["R_HE"], par["R_HN"],
                                         par["R_DD"], par["R_DF"], bono_antiguedad=2574)
    assert r["pago_horas_extra"] == D("1870.83")
    assert r["pago_nocturno"] == D("187.08")
    assert r["pago_domingos"] == D("2245.00")
    assert r["pago_feriados"] == D("1496.67")
    assert r["total_ganado"] == D("30823.58")


def test_rciva_sin_impuesto():
    par, _, _ = _datos()
    r = c.calcular_aportes_y_rciva(4500, 3300, 0, par)
    assert r["aportes_laborales"] == D("571.95")          # 12,71%
    assert r["impuesto_determinado"] == D("0.00")
    assert r["rciva_retener"] == D("0.00")
    assert r["saldo_favor_nuevo"] == D("0.00")             # el crédito de 1 SMN no genera saldo
    assert r["liquido_pagable"] == D("3928.05")


def test_rciva_con_impuesto():
    par, _, _ = _datos()
    r = c.calcular_aportes_y_rciva(15000, 3300, 1000, par)
    ans = D("23.00")                                       # 2.000 x 1,15%
    assert r["aporte_nacional_solidario"] == ans
    base = D("15000") - D("1906.50") - ans                 # AFP 12,71% = 1.906,50
    assert r["base_imponible"] == base
    imp = c.redondear((base - D("6600")) * D("0.13"))      # 841,165 -> 841,17
    assert r["impuesto_determinado"] == imp
    neto = imp - D("429.00")
    assert r["impuesto_neto"] == neto
    assert r["saldo_favor_fisco"] == neto - D("130.00")    # F-110: 13% de 1.000
    assert r["rciva_retener"] == neto - D("130.00")
    assert r["liquido_pagable"] == D("15000") - D("1906.50") - ans - r["rciva_retener"]


def test_rciva_contra_planilla_tributaria_de_referencia():
    """Filas de la planilla tributaria de ejemplo (SMN Bs 2.750)."""
    par = c.cargar_parametros({"SMN": 2750})
    casos = [
        # ingreso neto, F-110 (casilla 693), saldo anterior, factor UFV -> esperado
        (8097.55, 100, 100, "1.01", dict(exceso_imponible="2597.55", impuesto_determinado="337.68",
         credito_dependiente="357.50", impuesto_neto="0.00", saldo_favor_fisco="0.00",
         saldo_favor_dependiente="100.00", saldo_favor_anterior_actualizado="101.00",
         saldo_utilizado="0.00", rciva_retener="0.00", saldo_favor_nuevo="201.00")),
        (2500, 250, 0, "1", dict(exceso_imponible="0.00", impuesto_determinado="0.00",
         impuesto_neto="0.00", saldo_favor_dependiente="250.00", rciva_retener="0.00",
         saldo_favor_nuevo="250.00")),
        (2500, 0, 0, "1", dict(saldo_favor_nuevo="0.00", rciva_retener="0.00")),
        (18000, 100, 700, "1.0014285714285714285714285714", dict(
         exceso_imponible="12500.00", impuesto_determinado="1625.00", impuesto_neto="1267.50",
         saldo_favor_fisco="1167.50", saldo_favor_anterior_actualizado="701.00",
         saldo_utilizado="701.00", rciva_retener="466.50", saldo_favor_nuevo="0.00")),
    ]
    totales = {"impuesto_determinado": D(0), "rciva_retener": D(0), "saldo_favor_nuevo": D(0)}
    for ingreso, f110, ant, fac, esperado in casos:
        r = c.calcular_rciva(ingreso, 2750, D(str(f110)) / D("0.13"), par, ant, 1, D(fac))
        for k, v in esperado.items():
            assert r[k] == D(v), (ingreso, k, r[k], v)
        for k in totales:
            totales[k] += r[k]
    assert totales["impuesto_determinado"] == D("1962.68")   # total de la referencia
    assert totales["rciva_retener"] == D("466.50")
    assert totales["saldo_favor_nuevo"] == D("451.00")


def test_saldo_anterior_con_ufv():
    par, _, _ = _datos()
    r = c.calcular_aportes_y_rciva(4500, 3300, 0, par, saldo_favor_anterior=1000,
                                   ufv_anterior=2, ufv_actual=2.2)
    assert r["saldo_favor_anterior_actualizado"] == D("1100.00")
    assert r["mantenimiento_valor"] == D("100.00")
    assert r["saldo_favor_nuevo"] == D("1100.00")          # no se usó: no hay impuesto
    # si la UFV baja, no se reduce el saldo
    r = c.calcular_aportes_y_rciva(4500, 3300, 0, par, saldo_favor_anterior=1000,
                                   ufv_anterior=2, ufv_actual=1.9)
    assert r["saldo_favor_anterior_actualizado"] == D("1000.00")


def test_patronal_y_previsiones():
    par, _, _ = _datos()
    p = c.calcular_aportes_patronales(10000, par)
    assert p["aportes_patronales"] == D("1721.00")         # 17,21%
    pr = c.calcular_previsiones(10000, par)
    assert pr["aguinaldo"] == pr["indemnizacion"] == D("833.00")


def test_planilla_completa_y_errores():
    par, tabla, emp = _datos()
    emp[0]["Horas_Extra"] = 10
    emp.append({"ID_Empleado": "EMP-X", "Haber_Basico": 0, "Fecha_Ingreso": date(2020, 1, 1)})
    filas, errores = pl.calcular_todos(emp, par, tabla, 2026, 10)
    assert len(filas) == 5 and len(errores) == 1 and "EMP-X" in errores[0]
    for f in filas:  # identidad contable
        assert f["liquido_pagable"] == f["total_ganado"] - f["aportes_laborales"] - f["rciva_retener"]
    # 22.450 + bono 2.574 + 10 h extra (1.870,83)
    assert filas[0]["total_ganado"] == D("26894.83")


def _filas_demo():
    par, tabla, emp = _datos()
    emp[0].update(Departamento="Producción", Horas_Extra=10, Facturas_F110=5000)
    emp[1].update(Departamento="Producción", Dias_Domingo=1)
    return pl.calcular_todos(emp, par, tabla, 2026, 10, 3.2, 3.25)[0], par


def test_asientos_cuadran_y_respetan_departamentos():
    import asientos as a
    filas, _ = _filas_demo()
    for asiento in (a.asiento_privado(filas), a.asiento_publico(filas)):
        assert asiento["cuadra"] and asiento["total_debe"] > 0
        assert asiento["total_debe"] == sum(f["total_ganado"] + f["aportes_patronales"]
                                            + f["aguinaldo"] + f["indemnizacion"] for f in filas)
    solo_admin = [f for f in filas if f["departamento"] == a.ADMIN]
    txt = " ".join(l["cuenta"] for l in a.asiento_privado(solo_admin)["lineas"])
    assert "Producción" not in txt and "Administración" in txt
    codigos = {l["codigo"] for l in a.asiento_publico(solo_admin)["lineas"]}
    assert "1155" not in codigos and "611" in codigos
    solo_prod = [f for f in filas if f["departamento"] == a.PRODUCCION]
    codigos = {l["codigo"] for l in a.asiento_publico(solo_prod)["lineas"]}
    assert "1155" in codigos and "611" not in codigos


def test_asiento_privado_importes_clave():
    import asientos as a
    filas, _ = _filas_demo()
    g = lambda k: sum(f[k] for f in filas)
    lineas = {l["cuenta"]: l for l in a.asiento_privado(filas)["lineas"]}
    assert lineas["Sueldos y Salarios por pagar"]["haber"] == g("liquido_pagable")
    assert lineas["C.N.S. por pagar"]["haber"] == g("patronal_cns")
    assert lineas["RC-IVA dependientes por pagar"]["haber"] == g("rciva_retener")
    assert lineas["Aportes a la Gestora por Pagar"]["haber"] == (
        g("aportes_laborales") + g("aportes_patronales") - g("patronal_cns"))


def test_normalizar_departamento():
    import asientos as a
    assert a.normalizar_departamento("Producción") == a.PRODUCCION
    assert a.normalizar_departamento("PRODUCCION") == a.PRODUCCION
    assert a.normalizar_departamento("Comercializacion") == a.ADMIN
    assert a.normalizar_departamento(None) == a.ADMIN


def test_comprobantes_se_generan():
    import asientos as a
    import comprobantes as k
    from io import BytesIO
    from openpyxl import load_workbook
    filas, par = _filas_demo()
    meta = dict(empresa="Entidad Ejemplo", nit="", periodo="Octubre 2026",
                fecha_corte=date(2026, 10, 31), parametros=par, ufv_anterior=1, ufv_actual=1)
    ap, au = a.asiento_privado(filas), a.asiento_publico(filas)
    assert k.pdf_asiento_privado(ap, meta)[:5] == b"%PDF-"
    assert k.pdf_comprobante_publico(au, meta)[:5] == b"%PDF-"
    for xl in (k.excel_asiento_privado(ap, meta), k.excel_comprobante_publico(au, meta)):
        assert load_workbook(BytesIO(xl)).worksheets[0]["A1"].value == "Entidad Ejemplo"


def test_documentos_se_generan():
    from io import BytesIO
    from openpyxl import load_workbook
    par, tabla, emp = _datos()
    filas, _ = pl.calcular_todos(emp, par, tabla, 2026, 10, 3.2, 3.25)
    meta = dict(empresa="Empresa Ejemplo S.R.L.", nit="123", periodo="Octubre 2026",
                fecha_corte=date(2026, 10, 31), parametros=par, ufv_anterior=3.2, ufv_actual=3.25)
    for gen in (d.excel_planilla_laboral, d.excel_planilla_impositiva):
        wb = load_workbook(BytesIO(gen(filas, meta)))
        assert wb.worksheets[0]["A1"].value == "Empresa Ejemplo S.R.L."
    for gen in (d.pdf_planilla_laboral, d.pdf_planilla_impositiva):
        assert gen(filas, meta)[:5] == b"%PDF-"


if __name__ == "__main__":
    fallos = 0
    for nombre, fn in sorted(globals().items()):
        if nombre.startswith("test_") and callable(fn):
            try:
                fn()
                print("OK  ", nombre)
            except AssertionError as exc:
                fallos += 1
                print("FALLA", nombre, exc)
    sys.exit(1 if fallos else 0)
